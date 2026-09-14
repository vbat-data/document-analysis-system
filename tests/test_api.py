"""Tests for API endpoints."""


def test_root(client):
    """/ returns app info."""
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert "app" in data
    assert "version" in data


def test_health(client):
    """/health returns healthy."""
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert data["version"] == "0.1.0"
    # Note: ML model is lazy-loaded, so it isn't loaded in tests.
    assert data["models_loaded"] is False


def test_list_documents_empty(client):
    """Empty list on fresh DB."""
    response = client.get("/documents")
    assert response.status_code == 200
    assert response.json() == []


def test_get_missing_document(client):
    """404 for unknown document id."""
    response = client.get("/documents/doc_doesnotexist")
    assert response.status_code == 404
    assert "not found" in response.json()["detail"].lower()


def test_upload_unsupported_format(client):
    """400 for unsupported file format."""
    files = {"file": ("test.txt", b"hello world", "text/plain")}
    response = client.post("/documents/upload", files=files)
    assert response.status_code == 400
    assert "Unsupported format" in response.json()["detail"]


def test_upload_docx_success(client, monkeypatch):
    """Upload a real DOCX; ML is monkeypatched to return predictable entities."""
    from docx import Document
    from io import BytesIO

    # Stub the NER model so we don't download RuBERT during tests
    from src import ml_models

    class FakeModel:
        is_loaded = True

        def extract(self, text: str):
            return [
                {"entity_type": "ORG", "value": "Ромашка", "confidence": 0.95},
                {"entity_type": "DATE", "value": "01.01.2024", "confidence": 0.92},
            ]

    monkeypatch.setattr(ml_models, "get_ner_model", lambda: FakeModel())

    # Rebuild routes dependency: routes.py imported get_ner_model at import time
    from src.api import routes

    monkeypatch.setattr(routes, "get_ner_model", lambda: FakeModel())

    # Build DOCX in memory
    doc = Document()
    doc.add_paragraph("Договор поставки №42 от 01.01.2024")
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    files = {
        "file": (
            "contract.docx",
            buffer,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/documents/upload", files=files)
    assert response.status_code == 200, response.text

    data = response.json()
    assert data["document_type"] == "contract"
    assert data["status"] == "completed"
    assert data["validation"]["is_valid"] is True
    assert len(data["entities"]) == 2


def test_uploaded_document_appears_in_list(client, monkeypatch):
    """Uploaded doc is visible in the /documents listing."""
    from docx import Document
    from io import BytesIO

    from src.api import routes

    class FakeModel:
        is_loaded = True

        def extract(self, text: str):
            return [{"entity_type": "ORG", "value": "X", "confidence": 0.9}]

    monkeypatch.setattr(routes, "get_ner_model", lambda: FakeModel())

    doc = Document()
    doc.add_paragraph("Договор поставки №1")
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    client.post(
        "/documents/upload",
        files={"file": ("c.docx", buffer, "application/octet-stream")},
    )

    response = client.get("/documents")
    assert response.status_code == 200
    data = response.json()
    assert len(data) == 1
    assert data[0]["filename"] == "c.docx"
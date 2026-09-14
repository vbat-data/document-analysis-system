"""Tests for document processing logic."""
from pathlib import Path

import pytest

from src.document_processor import (
    classify_document,
    generate_document_id,
    parse_document,
)
from src.schemas import DocumentType


def test_generate_document_id_is_unique():
    """100 IDs should all be distinct."""
    ids = {generate_document_id() for _ in range(100)}
    assert len(ids) == 100


def test_generate_document_id_has_prefix():
    """ID starts with 'doc_'."""
    doc_id = generate_document_id()
    assert doc_id.startswith("doc_")
    assert len(doc_id) == 14  # "doc_" + 10 hex chars


@pytest.mark.parametrize(
    "text,expected",
    [
        ("Договор поставки №123", DocumentType.CONTRACT),
        ("Счёт на оплату товара", DocumentType.INVOICE),
        ("Акт выполненных работ", DocumentType.ACT),
        ("Произвольный текст без ключевых слов", DocumentType.UNKNOWN),
    ],
)
def test_classify_document(text: str, expected: DocumentType):
    """Document classification by keywords."""
    assert classify_document(text) == expected


def test_parse_document_with_unsupported_extension(tmp_path: Path):
    """Unsupported formats raise ValueError."""
    file_path = tmp_path / "note.txt"
    file_path.write_text("hello")

    with pytest.raises(ValueError, match="Unsupported file format"):
        parse_document(file_path)


def test_parse_docx(tmp_path: Path):
    """Parsing a real DOCX file works."""
    from docx import Document

    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Договор №42 от 01.01.2024")
    doc.add_paragraph("Поставщик: ООО Ромашка")
    doc.save(docx_path)

    text = parse_document(docx_path)
    assert "Договор" in text
    assert "Ромашка" in text
def test_extract_regex_dates():
    """Regex finds dates in common formats."""
    from src.document_processor import extract_regex_entities

    text = "Договор от 01.01.2024 и доп. соглашение от 15/05/2023"
    entities = extract_regex_entities(text)

    dates = [e.value for e in entities if e.entity_type == "DATE"]
    assert "01.01.2024" in dates
    assert "15/05/2023" in dates


def test_extract_regex_money():
    """Regex finds money amounts."""
    from src.document_processor import extract_regex_entities

    text = "Сумма договора: 1 500 000 руб. и 5000 рублей предоплата"
    entities = extract_regex_entities(text)

    money = [e.value for e in entities if e.entity_type == "MONEY"]
    assert any("1 500 000" in m for m in money)
    assert any("5000" in m for m in money)


def test_analyze_merges_ml_and_regex():
    """Analyze combines regex entities with ML entities."""
    from src.document_processor import analyze

    class FakeModel:
        def extract(self, text: str):
            return [
                {"entity_type": "ORG", "value": "ООО Ромашка", "confidence": 0.95},
            ]

    text = "Договор с ООО Ромашка от 01.01.2024 на сумму 150 000 руб."
    entities, doc_type, validation, elapsed = analyze(text, FakeModel())

    types = {e.entity_type for e in entities}
    assert "ORG" in types
    assert "DATE" in types
    assert "MONEY" in types
    assert validation.is_valid is True
    assert elapsed >= 0